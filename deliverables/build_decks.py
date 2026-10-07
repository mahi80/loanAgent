"""Builds the two client decks (10 slides each) from generated diagrams and
live prototype outputs.

Run:  python deliverables/diagrams.py && python deliverables/build_decks.py
"""
from __future__ import annotations

import os
import sys
from pathlib import Path

os.environ.setdefault("LLM_MODE", "mock")
from pptx import Presentation  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
from pptx.enum.text import PP_ALIGN  # noqa: E402
from pptx.util import Emu, Inches, Pt  # noqa: E402

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
IMG = HERE / "img"
sys.path.insert(0, str(ROOT))
from shared import vsm  # noqa: E402

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
TEAL = RGBColor(0x2A, 0x9D, 0x8F)
ORANGE = RGBColor(0xE7, 0x6F, 0x51)
GREY = RGBColor(0x6C, 0x75, 0x7D)
LIGHT = RGBColor(0xF1, 0xF5, 0xF9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x22, 0x22, 0x22)
W, H = Inches(13.333), Inches(7.5)


class Deck:
    def __init__(self, footer: str):
        self.p = Presentation()
        self.p.slide_width, self.p.slide_height = W, H
        self.footer = footer
        self.n = 0

    def slide(self, title: str, kicker: str = ""):
        s = self.p.slides.add_slide(self.p.slide_layouts[6])
        self.n += 1
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, Inches(0.12))
        bar.fill.solid()
        bar.fill.fore_color.rgb = TEAL
        bar.line.fill.background()
        if kicker:
            self.text(s, kicker.upper(), 0.5, 0.25, 12, 0.3, 11, TEAL, bold=True)
        self.text(s, title, 0.5, 0.5, 12.3, 0.8, 26, NAVY, bold=True)
        self.text(s, f"{self.footer}   |   {self.n}", 0.5, 7.05, 12.3, 0.3, 9, GREY)
        return s

    @staticmethod
    def text(s, txt, x, y, w, h, size=14, color=DARK, bold=False, align=PP_ALIGN.LEFT):
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = txt
        p.alignment = align
        p.font.size, p.font.bold, p.font.color.rgb = Pt(size), bold, color
        return tb

    @staticmethod
    def bullets(s, items, x, y, w, h, size=14, color=DARK):
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        for i, it in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            sub = it.startswith("  ")
            head, _, rest = it.strip().partition("|")
            r = p.add_run()
            r.text = ("– " if sub else "• ") + head
            r.font.size, r.font.color.rgb = Pt(size - (2 if sub else 0)), color
            r.font.bold = bool(rest) and not sub
            if rest:
                r2 = p.add_run()
                r2.text = rest
                r2.font.size, r2.font.color.rgb = Pt(size - (2 if sub else 0)), color
            p.level = 1 if sub else 0
            p.space_after = Pt(5)
        return tb

    @staticmethod
    def card(s, head, body, x, y, w, h, fill=LIGHT, head_color=NAVY, size=12):
        sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        sh.adjustments[0] = 0.08
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
        sh.line.color.rgb = head_color
        tf = sh.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.12)
        p = tf.paragraphs[0]
        p.text = head
        p.font.size, p.font.bold, p.font.color.rgb = Pt(size + 2), True, head_color
        for line in body.split("\n"):
            q = tf.add_paragraph()
            q.text = line
            q.font.size, q.font.color.rgb = Pt(size), DARK
        return sh

    @staticmethod
    def image(s, path, x, y, w=None, h=None):
        from PIL import Image

        iw, ih = Image.open(path).size
        if w and h:  # fit inside box keeping aspect ratio
            scale = min(w / iw, h / ih)
            w2, h2 = iw * scale, ih * scale
            x, y = x + (w - w2) / 2, y + (h - h2) / 2
            return s.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w2), Inches(h2))
        return s.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w) if w else None, Inches(h) if h else None)

    @staticmethod
    def table(s, rows, x, y, w, h, size=10, col_w=None, header_fill=NAVY):
        nr, nc = len(rows), len(rows[0])
        t = s.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(h)).table
        if col_w:
            for j, cw in enumerate(col_w):
                t.columns[j].width = Emu(int(Inches(w) * cw / sum(col_w)))
        for i, r in enumerate(rows):
            for j, v in enumerate(r):
                c = t.cell(i, j)
                c.text = str(v)
                c.margin_left = c.margin_right = Inches(0.05)
                c.margin_top = c.margin_bottom = Inches(0.02)
                for p in c.text_frame.paragraphs:
                    p.font.size = Pt(size)
                    p.font.bold = i == 0
                    p.font.color.rgb = WHITE if i == 0 else DARK
                c.fill.solid()
                c.fill.fore_color.rgb = header_fill if i == 0 else (LIGHT if i % 2 else WHITE)
        return t

    def kpis(self, s, items, y=1.45):
        n = len(items)
        w = (12.3 - (n - 1) * 0.2) / n
        for i, (big, small) in enumerate(items):
            x = 0.5 + i * (w + 0.2)
            sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(1.05))
            sh.fill.solid()
            sh.fill.fore_color.rgb = NAVY
            sh.line.fill.background()
            tf = sh.text_frame
            p = tf.paragraphs[0]
            p.text = big
            p.alignment = PP_ALIGN.CENTER
            p.font.size, p.font.bold, p.font.color.rgb = Pt(24), True, WHITE
            q = tf.add_paragraph()
            q.text = small
            q.alignment = PP_ALIGN.CENTER
            q.font.size, q.font.color.rgb = Pt(10), WHITE

    def notes(self, s, txt):
        s.notes_slide.notes_text_frame.text = txt

    def save(self, path):
        self.p.save(path)
        print("wrote", path.name, f"({min(self.n, 10)} main slides + {max(self.n - 10, 0)} appendix)")


def ba_rows(csv, demand):
    df = vsm.before_after(vsm.load(csv), vsm.Assumptions(demand))
    return [list(df.columns)] + df.astype(str).values.tolist()


def appendix_deployment(d: "Deck", asg: str) -> None:
    """Appendix (beyond the 10-slide main story): POC deployment, then AWS and Azure production targets."""
    s = d.slide("Appendix A · POC deployment today: Docker Compose with a local GPU LLM", "Appendix · deployment")
    d.image(s, IMG / "poc_deployment.png", 0.3, 1.25, 12.7, 5.75)
    d.notes(s, "What actually runs today. Same agents and code path as production; only the platform services change.")
    for letter, cloud, label in (("B", "aws", "AWS"), ("C", "azure", "Azure")):
        s = d.slide(f"Appendix {letter} · {label} production deployment: agents, MCP server, PostgreSQL",
                    "Appendix · deployment")
        d.image(s, IMG / f"{asg}_deploy_{cloud}.png", 0.3, 1.2, 12.7, 5.8)
        d.notes(s, f"Target {label} deployment. Client-cloud agnostic: the same container images, agent graph and "
                   "MCP tool contracts deploy to either cloud; only managed services differ.")


# ===================================================================== A1
def deck_a1():
    sys.path.insert(0, str(ROOT / "a1_loan_lifecycle"))
    from orchestrator import LoanOrchestrator, list_applications

    o = LoanOrchestrator()
    cases = [o.run_until_gate(a) for a in list_applications()]
    csv = ROOT / "a1_loan_lifecycle" / "data" / "vsm_steps.csv"
    a = vsm.Assumptions(120)
    cur = vsm.summary(vsm.state(vsm.load(csv), a), a)
    fut = vsm.summary(vsm.state(vsm.load(csv), a, True), a)

    d = Deck("Agentic Loan Lifecycle Management · FDE client challenge · synthetic data, mocked step times")

    s = d.slide("Faster credit decisions and earlier risk signals, with humans in control", "Executive summary")
    d.kpis(s, [(f"{cur['lead_time_d']:.0f} → {fut['lead_time_d']:.0f} d", "origination-to-disbursement lead time"),
               (f"{cur['pce']:.0%} → {fut['pce']:.0%}", "process cycle efficiency"),
               (f"{cur['nva_h']:.1f} → {fut['nva_h']:.1f} h", "manual NVA touch per file"),
               (f"{cur['capacity_per_month']} → {fut['capacity_per_month']}", "files / month at same FTEs"),
               ("0", "autonomous credit decisions")])
    d.bullets(s, [
        "Problem: |34.5-day lead time, only 6% of which is value-adding work; appraisal capacity (66 files/mo) is below demand (120), so backlog grows.",
        "Approach: |Lean value stream map → Automation Priority Index → 7 agents placed where waste and bottlenecks are, not everywhere.",
        "Prototype (working): |application → document ingestion → extraction with evidence → missing-info → RAG-cited policy checks → risk observations → recommendation → human approval → hash-chained decision summary.",
        "Control by design: |agents recommend, humans decide; authority matrix enforced in code; deterministic financial maths; prompt-injection and PII guards.",
        "Ask: |8-week POC in one segment (SME/mid-corporate), shadow mode on live files, MRM-ready evidence pack.",
    ], 0.5, 2.75, 12.3, 4.2, 15)
    d.notes(s, "All figures use mocked step times (stated) and synthetic data. KPIs are computed by shared/vsm.py.")

    s = d.slide("Credit teams spend most of the lifecycle waiting, chasing and re-keying", "Problem")
    d.image(s, IMG / "a1_lifecycle.png", 0.4, 1.35, 12.5, 3.6)
    d.card(s, "Manual effort", "Docs arrive by e-mail; data re-keyed into LOS and\nspreadsheets; memos written from scratch.", 0.5, 5.15, 4.0, 1.7)
    d.card(s, "Late risk visibility", "Policy checks on checklists; covenant breaches\nfound at quarterly test; adverse media ad hoc.", 4.65, 5.15, 4.0, 1.7)
    d.card(s, "Weak auditability", "Rationale spread across e-mail and memos;\nhard to evidence who decided what, on which data.", 8.8, 5.15, 4.0, 1.7)

    s = d.slide("Current process: value stream map shows 94% of elapsed time adds no value", "Current process · Lean VSM")
    d.image(s, IMG / "a1_vsm_block.png", 0.3, 1.25, 12.7, 5.75)
    d.notes(s, "Takt = 157.5 h / 120 applications = 1.31 h. PCE = VA / lead time. Step times mocked - validate in POC week 1 via time-and-motion + LOS timestamps.")

    s = d.slide("Opportunity: five areas where agents remove the most waste", "Opportunity · Automation Priority Index")
    d.image(s, IMG / "a1_api.png", 0.3, 1.3, 6.2, 3.8)
    pr = vsm.priority(vsm.load(csv), a).head(5)
    d.table(s, [["#", "Value area (VSM step)", "Agent", "API"]] +
            [[r.rank, r.step, r.agent, r.api] for r in pr.itertuples()], 6.7, 1.35, 6.2, 2.6, 10, [0.4, 2.4, 3.2, 0.6])
    d.bullets(s, [
        "Priority = waste hours × feasibility × bottleneck factor|, so effort goes where it moves lead time and capacity.",
        "Approval ranks low on purpose: |the decision stays human; agents only route and prepare the pack.",
        "Plus portfolio early-warning: |continuous covenant/EWS monitoring instead of quarterly tests.",
    ], 6.7, 4.25, 6.2, 2.6, 12)

    s = d.slide("Agent ecosystem: seven specialist agents, one orchestrator, one human gate", "Proposed AI / agent solution")
    d.image(s, IMG / "a1_agent_flow.png", 0.3, 1.25, 12.7, 3.2)
    d.table(s, [["Agent", "Does", "Never does"],
                ["Intake", "Normalises application, doc inventory, injection screen", "Accept unscreened docs into prompts"],
                ["Document Intelligence", "Extracts fields with confidence + evidence; grounding check", "Invent values not in source"],
                ["Due-Diligence", "Missing info + draft request; ratios; RAG-cited policy checks", "Compute ratios with the LLM"],
                ["Risk", "Transparent scorecard grade, observations, mitigants", "Use protected attributes"],
                ["Approval / Workflow", "Recommends next action, routes by authority, drafts memo", "Approve or decline"],
                ["Disbursement / CP", "Tracks conditions precedent; blocks release until evidenced", "Release funds"],
                ["Portfolio Monitoring", "Quarterly covenant + EWS scan, Red/Amber notes", "Change limits or ratings"]],
            0.5, 4.5, 12.3, 2.45, 10, [1.6, 5.5, 3.2])

    s = d.slide("Architecture: enterprise-ready, model-agnostic, auditable", "Architecture")
    d.image(s, IMG / "a1_architecture.png", 0.3, 1.2, 12.7, 5.8)

    s = d.slide("Demo: three applications, three different outcomes", "Working prototype · streamlit run a1_loan_lifecycle/app.py")
    rows = [["Case", "Facility", "Key agent findings", "Recommendation", "Routed to"]]
    for c in cases:
        ap, dd = c["approval"], c["dd"]
        fails = [f"{x['clause']} {x['check']} {x['value']}" for x in dd["checks"] if x["status"] in ("FAIL", "MISSING")]
        warns = [x["clause"] for x in dd["checks"] if x["status"] == "WARN"]
        miss = [m["item"] for m in dd["missing"]["items"]]
        find = "; ".join(fails[:3]) or "All checks pass"
        if miss:
            find += f" | missing: {', '.join(miss[:3])}"
        if warns:
            find += f" | warnings: {', '.join(warns)}"
        if c["intake"]["security_flags"]:
            find += " | prompt-injection text quarantined"
        rows.append([f"{c['id']}\n{c['application']['borrower']}", f"${c['application']['amount_usd'] / 1e6:.1f}M "
                     f"{c['application']['product'].replace('_', ' ')}", find, ap["recommendation"].replace("_", " "), ap["required_authority"]])
    d.table(s, rows, 0.5, 1.35, 12.3, 3.3, 10, [1.7, 1.4, 6.0, 1.6, 1.6])
    d.bullets(s, [
        "Tabs: |⓪ value stream · ① origination → decision (agent trace, evidence, cited policy, memo, human gate) · ② disbursement CPs · ③ portfolio EWS · ④ audit chain + LLM telemetry.",
        "Controls you can try: |approving above your authority is blocked; overriding the agent without rationale is blocked; every step lands in a hash-chained audit log.",
        "LLM: |Azure OpenAI when credentials are set (.env), deterministic mock otherwise, so the demo always runs; PII masked before any model call.",
    ], 0.5, 4.85, 12.3, 2.1, 12)

    s = d.slide("Business impact: before vs after (assumptions stated)", "Business impact")
    d.table(s, ba_rows(csv, 120), 0.5, 1.35, 6.3, 3.4, 11, [3, 1, 1, 1.4])
    d.table(s, [["KPI", "Before", "After", "Basis"],
                ["Document review effort / file", "8.5 h", "1.0 h", "VSM steps 2-3 NVA"],
                ["Exception identification", "sample-based", "100% of files", "every policy check runs on every file"],
                ["SLA adherence (21-day SLA)", "~30%", "~90%", "assumed lead-time distribution"],
                ["Analyst productivity", "1.0x", "~2.0x", "touch time 44 → 21 h"],
                ["Early warning lead", "at quarterly test", "+1 quarter", "continuous EWS scan"]],
            7.0, 1.35, 5.8, 3.0, 10, [2.4, 1.1, 1.1, 2.4])
    d.bullets(s, [
        "Value (illustrative): |1,440 files/yr × 22.75 h saved × $60/h loaded ≈ $2.0M/yr capacity; plus earlier interest accrual from a 22-day faster lead time.",
        "Assumptions: |120 applications/month, 21 days × 7.5 h, step times mocked (to be baselined in POC week 1); synthetic borrower data; 1 day = 8 h wait.",
    ], 0.5, 5.0, 12.3, 1.9, 12)

    s = d.slide("POC: 8 weeks to prove value in shadow mode", "POC approach")
    d.table(s, [["Weeks", "Milestone", "Exit criteria"],
                ["0-1", "Mobilise, data access, security review, VSM baseline (time-and-motion + LOS timestamps)", "Baseline KPIs signed off"],
                ["2-3", "Doc extraction + policy RAG on 200 historical files; evaluation harness", "≥95% field accuracy, ≥90% missing-info recall"],
                ["4-5", "Agents + orchestrator in LOS sandbox; analyst workbench; HITL + audit", "End-to-end on 50 files"],
                ["6-7", "Shadow mode on live SME/mid-corp files alongside analysts", "≥95% policy-check agreement, −40% touch time"],
                ["8", "Readout: KPI evidence, MRM/RAI pack, production plan", "Go / no-go for MVP"]],
            0.5, 1.35, 7.6, 3.3, 10, [0.7, 4.6, 2.5])
    d.card(s, "Team", "FDE lead · 2 AI engineers · data engineer · UX (0.5)\nClient: product owner, credit SME, risk/compliance,\nLOS integration engineer", 8.3, 1.35, 4.5, 1.55, size=11)
    d.card(s, "Data required", "200-500 historical files with decisions, credit policy,\nLOS extracts, KYC/sanctions samples, covenant data", 8.3, 3.05, 4.5, 1.55, size=11)
    d.card(s, "KPIs", "Field accuracy · missing-info recall · policy agreement ·\ntouch time · lead time · override rate · zero autonomous decisions", 0.5, 4.85, 6.1, 1.95, size=11)
    d.card(s, "Risks & mitigations", "Data access delays → masked extracts / synthetic\nHallucination → grounding check + deterministic maths\nMRM approval → early engagement, eval evidence\nAdoption → analyst co-design, override with reason", 6.75, 4.85, 6.05, 1.95, size=11, head_color=ORANGE)

    s = d.slide("Production roadmap: from one segment to an enterprise agent platform", "Production roadmap")
    for i, (h, b) in enumerate([
        ("POC · wks 1-8", "Shadow mode, SME/mid-corp\nEvidence pack for MRM\nBusiness case"),
        ("MVP · months 3-6", "Live in 1 segment, LOS integrated\nTeams approval cards\nSLA + override monitoring"),
        ("Scale · months 6-12", "Corporate & project finance\nDisbursement CP + portfolio EWS live\nAzure AI Search over memo history"),
        ("Platform · 12-18 months", "Reusable agent platform (gateway,\neval, audit, guardrails) for trade,\nretail credit, collections"),
    ]):
        d.card(s, h, b, 0.5 + i * 3.1, 1.4, 2.95, 2.5, size=12)
    d.bullets(s, [
        "Production hardening: |Entra ID + RBAC, private endpoints, Key Vault, data residency; APIM AI gateway with quotas and content safety; OpenTelemetry tracing.",
        "Model risk management: |model inventory, eval sets per agent, drift monitoring, challenger prompts, annual validation; documented human override.",
        "Next steps: |confirm POC segment and sponsor · grant data access · nominate credit SME · schedule week-0 VSM workshop.",
    ], 0.5, 4.1, 12.3, 1.5, 13)
    d.card(s, "Assumptions & synthetic data", "All borrowers, documents, policy (CP-x clauses), portfolio and step times are synthetic / mocked. "
           "Demand 120 apps/month; 21 days × 7.5 h; 1 wait-day = 8 h; $60/h loaded analyst cost. LLM runs in deterministic mock "
           "mode unless Azure OpenAI credentials are set. Ratios, rules and grades are deterministic code.",
           0.5, 5.65, 12.3, 1.25, fill=LIGHT, head_color=GREY, size=10)
    appendix_deployment(d, "a1")
    d.save(HERE / "A1_Loan_Lifecycle_Deck.pptx")


# ===================================================================== A2
def deck_a2():
    sys.path.insert(0, str(ROOT / "a2_chem_demand"))
    import pandas as pd
    from orchestrator import DemandResponseOrchestrator

    o = DemandResponseOrchestrator()
    o.sense()
    c = o.analyse(0)
    S = o.sensing
    csv = ROOT / "a2_chem_demand" / "data" / "vsm_steps.csv"
    a = vsm.Assumptions(120)
    cur = vsm.summary(vsm.state(vsm.load(csv), a), a)
    fut = vsm.summary(vsm.state(vsm.load(csv), a, True), a)
    opp = pd.read_csv(ROOT / "a2_chem_demand" / "data" / "opportunities.csv")
    opp["score"] = opp.value_score * opp.feasibility_score
    opp = opp.sort_values(["wave", "score"], ascending=[True, False])

    d = Deck("AI for Demand & Distribution · FDE client challenge · synthetic data, stated assumptions")

    s = d.slide("From forecasts to actions: demand sensing with an agentic response loop", "Executive summary")
    d.kpis(s, [(f"{cur['lead_time_d']:.0f} → {fut['lead_time_d']:.0f} d", "signal-to-action lead time"),
               (f"{S.mape_plan:.1%} → {S.mape_model:.1%}", "4-wk WMAPE (synthetic backtest)"),
               (f"{cur['pce']:.1%} → {fut['pce']:.0%}", "process cycle efficiency"),
               ("7", "use cases mapped & prioritised"),
               ("100%", "ERP actions human-approved")])
    d.bullets(s, [
        "Problem: |deviations from plan are found at the monthly S&OP, 3-4 weeks late; planners spend 70% of their time collecting and reconciling data.",
        "Opportunity map: |7 use cases scored on value × feasibility; demand sensing + agentic response leads wave 1 and is the data foundation for inventory optimisation and stockout prediction.",
        "Prototype (working): |orders + inventory + capacity + PMI/Brent/construction → sensing model → Deviation → Impact → Constraint → Recommender agents → planner approval → SAP STO / planned order / IBP update (mocked).",
        "Ask: |8-week POC on one product family across two regions with real SAP/IBP extracts, run in shadow mode next to the S&OP process.",
    ], 0.5, 2.75, 12.3, 4.2, 15)

    s = d.slide("Current process: deviations wait for the monthly S&OP (PCE 3.7%)", "Problem & current process · Lean VSM")
    d.image(s, IMG / "a2_vsm_block.png", 0.3, 1.25, 12.7, 5.75)

    s = d.slide("Use-case opportunity map: 7 opportunities, each tied to a KPI and an owner", "Opportunity map")
    d.table(s, [["Use case", "Business problem", "AI / agent approach", "Data required", "Action → human involvement",
                 "KPI · expected benefit (assumption)"]] +
            [[f"{r.id} {r.use_case} (wave {r.wave})", r.business_problem, r.ai_agent_approach, r.data_required,
              f"{r.action_taken} → {r.human_involvement}", f"{r.kpi} · {r.expected_benefit}"] for r in opp.itertuples()],
            0.3, 1.25, 12.75, 5.7, 7, [1.3, 2.0, 2.2, 2.0, 2.6, 2.2])
    d.notes(s, "Data required per use case is in a2_chem_demand/data/opportunities.csv and the app's Opportunity map tab.")

    s = d.slide("Prioritisation: demand sensing first - fixes the #1 bottleneck", "Prioritisation")
    d.image(s, IMG / "a2_opportunity_matrix.png", 0.3, 1.3, 5.6, 4.6)
    d.image(s, IMG / "a2_api.png", 6.0, 1.3, 6.9, 3.9)
    d.bullets(s, [
        "Wave 1 (POC): |UC1 demand sensing + agentic response, UC2 stockout/excess prediction (same data).",
        "Wave 2: |UC3 inventory optimisation, UC4 shipment exceptions, UC7 invoice automation (quick win).",
        "Wave 3: |UC5 network/lane optimisation, UC6 supply disruption prediction.",
    ], 6.0, 5.3, 6.9, 1.7, 12)

    s = d.slide("Solution: an agent loop that turns a demand signal into an approved action", "Proposed AI / agent solution")
    d.image(s, IMG / "a2_agent_flow.png", 0.3, 1.25, 12.7, 3.2)
    d.table(s, [["Agent", "Does", "Guardrail"],
                ["Deviation", "Flags SKU×region where sensed demand vs plan > 15% (or actuals > 20%, |z| > 1.5); ranks by margin at stake", "Thresholds configurable, logged"],
                ["Impact", "Drills to customers driving the change; flags anomalies (e.g. -86% outage)", "Facts only; narrative from data"],
                ["Constraint", "DC stock, in-transit, safety stock, other DCs' surplus, plant free capacity, lanes", "Deterministic maths"],
                ["Recommender", "Costs transfer / pull-forward / expedite / allocation; least-cost plan", "LLM explains, never sets qty"],
                ["Executor", "STO, planned order, aATP allocation, IBP key figure, notifications", "Only after approval"]],
            0.5, 4.55, 12.3, 2.4, 10, [1.3, 7.0, 2.6])

    s = d.slide("Architecture: SAP-integrated, ML + LLM, human-approved writes", "Architecture")
    d.image(s, IMG / "a2_architecture.png", 0.3, 1.2, 12.7, 5.8)

    s = d.slide("Demo: EPX-200 epoxy resin in APAC, from signal to SAP in one flow", "Working prototype · streamlit run a2_chem_demand/app.py")
    d.image(s, IMG / "a2_forecast.png", 0.3, 1.25, 6.5, 2.9)
    d.image(s, IMG / "a2_impact_plan.png", 6.8, 1.25, 6.2, 2.9)
    t = c["recommendation"]["totals"]
    d.bullets(s, [
        f"Detect: |sensed next-4-week demand {c['alert']['deviation_pct']:+.0%} vs frozen plan; actuals last 4 weeks {c['alert']['actual_dev_pct']:+.0%} (z={c['alert']['z_score']:+.1f}).",
        f"Impact: |{c['impact']['primary_driver']} drives {c['impact']['driver_share']:.0%} of the change; revenue exposure ${c['impact']['revenue_at_risk_usd']:,.0f}.",
        f"Constraints: |DC-Singapore covers {c['constraints']['position']['weeks_cover']} weeks; gap {t['need_t']:,} t vs safety stock; Rotterdam has surplus, Jurong has free capacity.",
        f"Recommend: |transfer + pull-forward + expedite + allocation: cost ${t['cost_usd']:,.0f} vs ${t['benefit_usd']:,.0f} margin protected.",
        f"Approve & act: |routed to {c['recommendation']['approval_level']} (> $50k); a planner without authority is blocked; approval creates SAP STO, planned order and IBP update (mocked) plus a notification.",
    ], 0.5, 4.3, 12.3, 2.7, 12)

    s = d.slide("Business impact: before vs after (assumptions stated)", "Business impact")
    d.table(s, ba_rows(csv, 120), 0.5, 1.35, 6.3, 3.4, 11, [3, 1, 1, 1.4])
    d.table(s, [["Lever", "Impact (assumption)", "Annual value"],
                ["Forecast error (4-wk WMAPE)", f"{S.mape_plan:.1%} → {S.mape_model:.1%} synthetic; 35% → 22% client assumption", "enabler"],
                ["Stockouts avoided", "−30% lost-sales incidents", "$6-9M"],
                ["Excess & obsolete", "−15% E&O", "$3-5M"],
                ["Expedite / premium freight", "−20%", "$4-6M"],
                ["Planner productivity", "NVA 20.5 → 3.3 h per exception", "~6 FTE redeployed"]],
            7.0, 1.35, 5.8, 3.1, 10, [2.2, 2.6, 1.0])
    d.bullets(s, [
        "Assumptions: |$5B revenue manufacturer, ~120 significant demand exceptions/month, 12 SKUs × 4 regions in the prototype; step times mocked; benefits ranges to be validated with the client's SAP history in POC weeks 1-2.",
        "Measured in prototype: |backtest WMAPE over 16 weekly origins on synthetic data (sensing vs frozen plan vs 13-week moving average).",
    ], 0.5, 5.0, 12.3, 1.9, 12)

    s = d.slide("POC: 8 weeks, one product family, two regions, shadow mode", "POC approach")
    d.table(s, [["Weeks", "Milestone", "Exit criteria"],
                ["0-1", "Mobilise; SAP/IBP extracts (3 yrs orders, stock, capacity); VSM baseline with planners", "Data contract + baseline KPIs"],
                ["2-3", "Feature pipeline + external signals; sensing model + backtest vs current plan", "≥20% WMAPE improvement"],
                ["4-5", "Agents + planner cockpit + Teams approval; SAP sandbox (STO / planned order)", "End-to-end on 10 historical events"],
                ["6-7", "Shadow mode next to weekly planning; capture accept/edit/reject", "≥60% recommendations accepted"],
                ["8", "Readout: value case, scale plan to all families/regions", "Go / no-go"]],
            0.5, 1.35, 7.6, 3.3, 10, [0.7, 4.6, 2.5])
    d.card(s, "Team", "FDE lead · data scientist · 2 AI/data engineers ·\nSAP integration (0.5) · client: demand planning lead,\nS&OP owner, IT/SAP basis", 8.3, 1.35, 4.5, 1.55, size=11)
    d.card(s, "Data required", "Orders/shipments, inventory by DC, production plans &\ncapacity, lanes & freight rates, promos, PMI/feedstock", 8.3, 3.05, 4.5, 1.55, size=11)
    d.card(s, "KPIs", "WMAPE & bias · detection lag · alert precision ·\nacceptance rate · OTIF · expedite cost · planner hours", 0.5, 4.85, 6.1, 1.95, size=11)
    d.card(s, "Risks & mitigations", "Data quality / master data → profiling in week 1\nSAP access → sandbox + mocked APIs first\nPlanner trust → explain every alert, edit before approve\nScope creep → one family, two regions", 6.75, 4.85, 6.05, 1.95, size=11, head_color=ORANGE)

    s = d.slide("Production roadmap: from one loop to an autonomous-with-oversight supply chain", "Production roadmap")
    for i, (h, b) in enumerate([
        ("POC · wks 1-8", "Epoxy family × APAC/EU\nShadow mode, value case"),
        ("MVP · months 3-6", "All families in 2 regions\nLive SAP writes after approval\nUC2 stockout/excess"),
        ("Scale · months 6-12", "Global; UC3 inventory optimisation\nUC4 shipment exceptions\nUC7 invoice automation"),
        ("Transform · 12-24 months", "UC5 network optimisation, UC6\nsupply risk; auto-approve low-value\nactions within policy"),
    ]):
        d.card(s, h, b, 0.5 + i * 3.1, 1.4, 2.95, 2.5, size=12)
    d.bullets(s, [
        "Platform: |shared feature store + MLflow, agent orchestrator, approval and audit services reused by every use case.",
        "Governance: |action-value thresholds by role, weekly model performance review, audit trail of every ERP write.",
        "Next steps: |confirm POC scope and sponsor · SAP/IBP extract request · planner workshop for the VSM baseline.",
    ], 0.5, 4.1, 12.3, 1.5, 13)
    d.card(s, "Assumptions & synthetic data", "104 weeks of synthetic orders (12 SKUs × 4 regions × 12 customers), PMI/Brent/construction "
           "signals, DC inventory, plant capacity and lanes, with two injected events (APAC epoxy project ramp; EU PVC outage). "
           "VSM step times mocked; financial benefits are ranges for a ~$5B manufacturer. SAP/IBP calls are mocked payloads.",
           0.5, 5.65, 12.3, 1.25, fill=LIGHT, head_color=GREY, size=10)
    appendix_deployment(d, "a2")
    d.save(HERE / "A2_Chemical_Demand_Deck.pptx")


if __name__ == "__main__":
    # each assignment has its own `orchestrator` / `agents` modules -> build in separate processes
    if len(sys.argv) > 1:
        {"a1": deck_a1, "a2": deck_a2}[sys.argv[1]]()
    else:
        import subprocess

        for which in ("a1", "a2"):
            subprocess.run([sys.executable, __file__, which], check=True)
